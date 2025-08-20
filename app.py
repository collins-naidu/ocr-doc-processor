import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import sv_ttk
import json
import os
import pytesseract
from PIL import Image
import docx
import fitz  # PyMuPDF

# You might need to configure the path to the Tesseract executable
# pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Document Converter")
        self.geometry("600x400")

        self.check_tesseract()

        self.config_file = "config.json"
        self.target_folder = self.load_target_folder()

        # Menu
        self.menu = tk.Menu(self)
        self.config(menu=self.menu)
        self.file_menu = tk.Menu(self.menu, tearoff=False)
        self.menu.add_cascade(label="File", menu=self.file_menu)
        self.file_menu.add_command(label="Set Target Folder", command=self.set_target_folder)
        self.file_menu.add_separator()
        self.file_menu.add_command(label="Exit", command=self.quit)

        # Main frame
        self.main_frame = ttk.Frame(self, padding="10")
        self.main_frame.pack(fill=tk.BOTH, expand=True)

        # Label for target folder
        self.target_folder_label = ttk.Label(self.main_frame, text=f"Target Folder: {self.target_folder}")
        self.target_folder_label.pack(pady=5)

        # Upload button
        self.upload_button = ttk.Button(self.main_frame, text="Upload Files", command=self.upload_files)
        self.upload_button.pack(pady=20)

        # Progress label
        self.progress_label = ttk.Label(self.main_frame, text="")
        self.progress_label.pack(pady=5)


        sv_ttk.set_theme("light")

    def load_target_folder(self):
        if os.path.exists(self.config_file):
            with open(self.config_file, 'r') as f:
                config = json.load(f)
                return config.get("target_folder", "")
        return ""

    def save_target_folder(self, folder):
        with open(self.config_file, 'w') as f:
            json.dump({"target_folder": folder}, f)
        self.target_folder = folder
        self.target_folder_label.config(text=f"Target Folder: {self.target_folder}")


    def set_target_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.save_target_folder(folder)
            messagebox.showinfo("Success", f"Target folder set to: {folder}")

    def upload_files(self):
        if not self.target_folder:
            messagebox.showerror("Error", "Please set a target folder first.")
            return

        files = filedialog.askopenfilenames(
            title="Select files",
            filetypes=(
                ("PDF files", "*.pdf"),
                ("Image files", "*.png *.jpg *.jpeg *.bmp *.tiff"),
                ("All files", "*.*")
            )
        )

        if files:
            self.process_files(files)

    def show_wait_message(self):
        self.progress_label.config(text="Processing files, please wait...")
        self.update_idletasks()

    def hide_wait_message(self):
        self.progress_label.config(text="")
        self.update_idletasks()


    def process_files(self, files):
        self.show_wait_message()
        try:
            for file_path in files:
                filename = os.path.basename(file_path)
                # Determine file type and process
                if filename.lower().endswith('.pdf'):
                    self.process_pdf(file_path)
                elif filename.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.tiff')):
                    self.process_image(file_path)
                else:
                    messagebox.showwarning("Unsupported File", f"File type for {filename} is not supported.")
            messagebox.showinfo("Success", "All files processed successfully.")
        except Exception as e:
            messagebox.showerror("Error", f"An error occurred: {e}")
        finally:
            self.hide_wait_message()

    def process_pdf(self, file_path):
        try:
            doc = fitz.open(file_path)
            text = ""
            is_ocr_needed = False

            # First, try to extract text directly
            for page in doc:
                page_text = page.get_text()
                if page_text.strip():  # If there is text
                    text += page_text

            # If no text was extracted, we need to do OCR
            if not text.strip():
                is_ocr_needed = True
                for page_num, page in enumerate(doc):
                    print(f"Page {page_num+1} of {file_path} requires OCR.")
                    pix = page.get_pixmap()
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    text += pytesseract.image_to_string(img)

            # Create and save the docx file
            docx_doc = docx.Document()
            docx_doc.add_paragraph(text)

            base_filename = os.path.splitext(os.path.basename(file_path))[0]
            output_filename = os.path.join(self.target_folder, f"{base_filename}.docx")

            docx_doc.save(output_filename)
            print(f"Successfully converted {file_path} to {output_filename}. OCR needed: {is_ocr_needed}")

        except Exception as e:
            messagebox.showerror("PDF Processing Error", f"Failed to process {file_path}: {e}")

    def check_tesseract(self):
        try:
            pytesseract.get_tesseract_version()
        except pytesseract.TesseractNotFoundError:
            messagebox.showerror("Tesseract Not Found",
                                 "Tesseract is not installed or not in your PATH. "
                                 "Please install Tesseract and try again.")
            self.quit()

    def process_image(self, file_path):
        try:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)

            doc = docx.Document()
            doc.add_paragraph(text)

            base_filename = os.path.splitext(os.path.basename(file_path))[0]
            output_filename = os.path.join(self.target_folder, f"{base_filename}.docx")

            doc.save(output_filename)
            print(f"Successfully converted {file_path} to {output_filename}")

        except Exception as e:
            messagebox.showerror("Image Processing Error", f"Failed to process {file_path}: {e}")


if __name__ == "__main__":
    app = App()
    app.mainloop()
